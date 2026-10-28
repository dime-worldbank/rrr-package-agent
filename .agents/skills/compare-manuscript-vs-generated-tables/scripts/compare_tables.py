#!/usr/bin/env python3
"""Compare manuscript tables (LaTeX) to generated output tables.

This script performs a numeric first-pass comparison and writes:
- Markdown verification report
- JSON diagnostics

Comparison is gated by Step 2 run status. If Step 2 failed, the script writes
an explicit failure report and does not attempt table comparison.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, List, Optional, Tuple


try:
    from openpyxl import load_workbook
except Exception as exc:  # pragma: no cover
    print("ERROR: openpyxl is required. Install with: pip install openpyxl", file=sys.stderr)
    raise

try:
    import fitz  # PyMuPDF
except ImportError:  # pragma: no cover
    # PDF text extraction is a nice-to-have, not a hard requirement: without
    # it, generated .pdf outputs are simply routed to the visual-review
    # bucket for manual comparison instead of being auto-verified, same as
    # before this feature existed. `pip install pymupdf` to enable it.
    fitz = None


NUMBER_RE = re.compile(r"[-+]?\d[\d,]*(?:\.\d+)?")
FAIL_RE = re.compile(
    r"\b(error|failed|failure|traceback|exception|file not found|no such file)\b",
    flags=re.IGNORECASE,
)
# Step 2 run logs are often reviewer-written prose (e.g. run_log.md), not raw
# console output, and routinely say things like "loaded without error" or
# "no errors found" - a plain keyword match on FAIL_RE reads that as a failure.
# Treat a match as a false positive if immediately preceded by a negation cue.
NEGATION_WORDS = {"without", "no", "zero", "any", "not"}


@dataclass
class PaperTable:
    index: int
    caption: str
    numeric_values: List[float]
    row_labels: List[str]
    label_token: Optional[str] = None


@dataclass
class OutputTable:
    index: int
    path: Path
    numeric_values: List[float]
    row_labels: List[str]


def detect_step2_failure(step2_log: Path) -> Tuple[bool, str]:
    if not step2_log.exists():
        return False, ""

    lines = read_text(step2_log).splitlines()
    for raw in lines:
        line = " ".join(raw.split())
        if not line:
            continue
        for match in FAIL_RE.finditer(line):
            preceding_words = line[: match.start()].split()
            nearby = {w.lower().strip(",.:;()") for w in preceding_words[-3:]}
            if nearby & NEGATION_WORDS:
                continue
            return True, line
    return False, ""


def natural_sort_key(path: Path) -> Tuple:
    parts = re.split(r"(\d+)", path.name)
    key = []
    for part in parts:
        if part.isdigit():
            key.append(int(part))
        else:
            key.append(part.lower())
    return tuple(key)


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def strip_latex_commands(text: str) -> str:
    text = re.sub(r"%.*", "", text)
    # esttab/estout emit row breaks like "\\[-1.8ex]" to add vertical space.
    # The generic "\command[opt]{arg}" stripper below requires a letter right
    # after the backslash, so it never matches this - the "-1.8" inside the
    # brackets would otherwise leak into extract_numbers() as a fake data
    # point and throw off every downstream positional comparison.
    text = re.sub(r"\\\\\s*\[[^\]]*\]", " ", text)
    text = re.sub(r"\\multicolumn\{[^}]*\}\{[^}]*\}\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\multirow\{[^}]*\}\{[^}]*\}\{([^}]*)\}", r"\1", text)
    # \cmidrule(lr){2-3} uses parentheses for its trim spec, which the generic
    # "\command[opt]{arg}" stripper below doesn't recognize (it only handles a
    # bracketed optional arg) - it leaves "(lr){2-3}" behind, and "2-3" then
    # reads as the two numbers 2 and -3 to extract_numbers(). Grouped-header
    # tables (any results table with column sub-groups) all use this command.
    text = re.sub(r"\\cmidrule(\([^)]*\))?\{[^}]*\}", " ", text)
    text = re.sub(r"\\textbf\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\emph\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\textit\{([^}]*)\}", r"\1", text)
    text = text.replace("\\quad", " ")
    text = text.replace("\\,", "")
    text = text.replace("$", "")
    text = re.sub(r"\\[a-zA-Z]+\*?(\[[^\]]*\])?(\{[^}]*\})?", " ", text)
    text = text.replace("{", " ").replace("}", " ")
    return " ".join(text.split())


COLUMN_INDEX_RE = re.compile(r"\(\s*\d{1,3}\s*\)")


def extract_numbers(text: str) -> List[float]:
    # Column-header index markers - "(1)", "(2)", "(3)" - are a near-universal
    # econ-table convention for labeling columns, not data. Standard errors in
    # parens always carry a decimal point ("(0.034)"), so a bare integer in
    # parens is never a real result; leaving these in would let the manuscript
    # and a generated output disagree purely on how many columns each one's
    # header row happens to enumerate, which then misaligns every real data
    # point that follows under the positional match below.
    text = COLUMN_INDEX_RE.sub(" ", text)
    values: List[float] = []
    for raw in NUMBER_RE.findall(text):
        cleaned = raw.replace(",", "")
        try:
            values.append(float(cleaned))
        except ValueError:
            continue
    return values


def normalize_label(text: str) -> str:
    lowered = text.lower()
    lowered = re.sub(r"[^a-z0-9 ]+", " ", lowered)
    lowered = " ".join(lowered.split())
    return lowered


def canonicalize_label(text: str) -> str:
    n = normalize_label(text)

    # Check the space-collapsed form for "R2 Adj" *before* the standalone-
    # digit strip below. Manuscript captions render R^2 with a caret, which
    # normalize_label turns into a space ("R$^2$ Adj." -> "r 2 adj"), while
    # code-generated tables usually just write "R2 Adj." (-> "r2 adj") - the
    # digit-strip only removes a digit that is its own token (bounded by
    # non-word chars on both sides) and silently no-ops on "r2" since there's
    # no boundary between the letter and the digit, so the two forms never
    # converge without this check.
    if n.replace(" ", "") == "r2adj":
        return "adjusted r2"

    n = re.sub(r"\b\d+\b", " ", n)
    n = " ".join(n.split())

    if "trade hospit" in n:
        return "trade and hospitality"

    if n.startswith("number of "):
        n = n.replace("number of ", "", 1)

    if "adjusted r" in n:
        return "adjusted r2"
    if "within r" in n:
        return "within r2"
    if n in {"r", "r2", "r 2"} or n.startswith("r 2"):
        return "r2"

    if "interaction" in n or n.startswith("interact"):
        if "placebo" in n or "energy" in n:
            return "interact placebo"
        return "interact"

    if "country sector" in n:
        return "country sector fe"
    if "country year" in n:
        return "country year fe"
    if "cyber incidents" in n:
        return "cyber incidents"
    if "countries" in n:
        return "countries"
    if "share 2014" in n:
        return "share 2014"

    return n


def parse_tabular_rows(raw_tabular_body: str) -> List[str]:
    row_labels: List[str] = []
    rows = raw_tabular_body.split("\\\\")
    for row in rows:
        row = row.strip()
        # A leading "[-1.8ex]"-style spacing argument is what's left over
        # after splitting on the "\\" row separator that preceded it in the
        # source (see strip_latex_commands for the same pattern pre-split).
        row = re.sub(r"^\[[^\]]*\]\s*", "", row)
        if not row:
            continue
        if "\\hline" in row or "\\midrule" in row or "\\toprule" in row or "\\bottomrule" in row:
            continue
        if "\\cline" in row or "\\cmidrule" in row:
            continue
        if row.startswith("\\multicolumn"):
            continue

        first_cell = row.split("&", 1)[0]
        cleaned = strip_latex_commands(first_cell).strip()
        if not cleaned:
            continue

        normalized = normalize_label(cleaned)
        if not normalized:
            continue

        # Skip likely headers and note lines.
        if normalized.startswith("note"):
            continue
        if "dependent variable" in normalized:
            continue
        if normalized in {
            "model",
            "variables",
            "sample structure",
            "fixed effects",
            "fit statistics",
            "vulnerability type",
            "all countries",
            "hics",
            "non hics",
            "hic",
            "non hic",
        }:
            continue

        row_labels.append(cleaned)
    return row_labels


def parse_paper_tables(tex_path: Path) -> List[PaperTable]:
    content = read_text(tex_path)
    table_blocks = list(
        re.finditer(r"\\begin\{table\}.*?\\end\{table\}", content, flags=re.S)
    )
    parsed: List[PaperTable] = []

    for idx, block in enumerate(table_blocks, start=1):
        block_text = block.group(0)
        cap_match = re.search(r"\\caption\{(.*?)\}", block_text, flags=re.S)
        caption = cap_match.group(1).strip() if cap_match else f"Table {idx}"

        # \label{tab:N} / \label{tab:A1} is the authoritative table number -
        # most manuscript captions are prose ("Summary Statistics") and never
        # literally spell out "Table N", so the caption text alone can't be
        # used to identify which exhibit this is.
        label_match = re.search(r"\\label\{tab:([A-Za-z]?\d+[A-Za-z]?)\}", block_text)
        label_token = label_match.group(1).upper() if label_match else extract_table_token(caption)

        tab_match = re.search(
            r"\\begin\{tabular\*?\}\{.*?\}(.*?)\\end\{tabular\*?\}",
            block_text,
            flags=re.S,
        )
        body = tab_match.group(1) if tab_match else block_text

        # Exclude the column-header row (before \midrule) from numeric
        # extraction. Simple "(1)" index headers are already stripped by
        # COLUMN_INDEX_RE, but descriptive headers like "Micro (1--9) & Small
        # (10--19) & ... & Large (50+)" contain real digits as prose, not
        # data - and the en-dash range syntax ("1--9") even fools the number
        # regex into reading the dash as a minus sign. Left in, these numbers
        # occupy the front of the value list and shift every real data point
        # that follows out of position against the output's values.
        midrule_match = re.search(r"\\midrule", body)
        data_body = body[midrule_match.end():] if midrule_match else body

        cleaned = strip_latex_commands(data_body)
        nums = extract_numbers(cleaned)
        row_labels = parse_tabular_rows(body)
        parsed.append(
            PaperTable(
                index=idx,
                caption=caption,
                numeric_values=nums,
                row_labels=row_labels,
                label_token=label_token,
            )
        )

    return parsed


def parse_output_workbook(path: Path) -> List[float]:
    values: List[float] = []
    wb = load_workbook(path, read_only=True, data_only=True)
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                if cell is None:
                    continue
                if isinstance(cell, (int, float)):
                    values.append(float(cell))
                elif isinstance(cell, str):
                    values.extend(extract_numbers(cell))
    wb.close()
    return values


def parse_output_row_labels(path: Path) -> List[str]:
    labels: List[str] = []
    wb = load_workbook(path, read_only=True, data_only=True)
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            if not row:
                continue
            for cell in row:
                if not isinstance(cell, str):
                    continue
                label = " ".join(cell.split())
                norm = normalize_label(label)
                if not norm:
                    continue
                if NUMBER_RE.fullmatch(label.replace(",", "")):
                    continue
                labels.append(label)
    wb.close()
    # Deduplicate while preserving order.
    deduped = list(dict.fromkeys(labels))
    return deduped


NUMERIC_TABLE_EXTS = (
    "*.xlsx", "*.txt", "*.csv", "*.tsv", "*.tex",
    # Stata/R export tables in these formats too (outreg2 .rtf, stargazer
    # type="html", Stata .log) - without these, files in these formats would
    # match neither NUMERIC_TABLE_EXTS nor VISUAL_REVIEW_EXTS and silently
    # disappear from the report entirely.
    "*.log", "*.html", "*.htm", "*.rtf",
)
# .pdf is deliberately not in either static list - it's classified dynamically
# at runtime (see extract_pdf_table): a PDF with a real text layer (e.g. R's
# gt::gtsave() via a headless-browser print, common for these packages) is
# numerically compared like any other table; one with no usable text (a
# scanned/rasterized PDF) falls back to the visual-review bucket.
VISUAL_REVIEW_EXTS = ("*.png", "*.jpg", "*.jpeg")
PDF_GLOB = "*.pdf"
PDF_MINUS_RE = re.compile("−")  # U+2212 MINUS SIGN, not ASCII "-"


def _is_pure_numeric_line(line: str) -> bool:
    """True if a line is only number token(s) and whitespace, nothing else.

    PDF text extraction lays a table's cells out one-per-line (or a few
    adjacent cells sharing a line), but header/title text can embed real
    digits too (e.g. a "Duration2011" column header) - blindly scanning every
    line for numbers would count those as data. Restricting extraction to
    lines that are *purely* numeric avoids that without assuming a fixed
    one-cell-per-line layout (which doesn't hold across every gt/webshot2
    table export).
    """
    if not NUMBER_RE.search(line):
        return False
    return not NUMBER_RE.sub("", line).strip()


def extract_pdf_table(path: Path) -> Optional[Tuple[List[float], List[str]]]:
    """Extract a table's values/labels directly from a PDF's text layer.

    Returns None (caller falls back to the visual-review bucket) if PyMuPDF
    isn't installed, the file can't be opened, or it yields negligible/no
    text - a strong sign the PDF is a scanned image rather than a real
    text-layer export, which this can't help with.
    """
    if fitz is None:
        return None
    try:
        doc = fitz.open(path)
        text = "\n".join(page.get_text() for page in doc)
        doc.close()
    except Exception:
        return None

    # R's gt package renders the minus sign as U+2212, not ASCII "-"; left
    # unnormalized, a negative value's sign is silently dropped rather than
    # the value being read correctly (and the line then fails the
    # pure-numeric check above because "−" isn't part of NUMBER_RE).
    text = PDF_MINUS_RE.sub("-", text)

    if len(text.strip()) < 20:
        return None

    values: List[float] = []
    labels: List[str] = []
    for raw in text.splitlines():
        line = " ".join(raw.split())
        if not line:
            continue
        if _is_pure_numeric_line(line):
            values.extend(extract_numbers(line))
        else:
            labels.append(line)

    if not values:
        return None

    deduped_labels = list(dict.fromkeys(labels))
    return values, deduped_labels


def parse_output_tex(path: Path) -> Tuple[List[float], List[str]]:
    """Parse a standalone generated .tex table (e.g. from esttab/outreg2/estout).

    Package-generated tables are not always .xlsx/.csv/.txt - some Stata/R
    workflows export .tex directly. Reuse the same LaTeX stripping/parsing
    logic used for the manuscript tables so these compare on equal footing.
    """
    content = read_text(path)
    tab_match = re.search(
        r"\\begin\{tabular\*?\}\{.*?\}(.*?)\\end\{tabular\*?\}", content, flags=re.S
    )
    body = tab_match.group(1) if tab_match else content
    cleaned = strip_latex_commands(body)
    nums = extract_numbers(cleaned)
    row_labels = parse_tabular_rows(body)
    return nums, row_labels


def strip_html_tags(text: str) -> str:
    text = re.sub(r"<!--.*?-->", " ", text, flags=re.S)
    text = re.sub(r"<(script|style)\b[^>]*>.*?</\1>", " ", text, flags=re.S | re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    text = text.replace("&nbsp;", " ").replace("&amp;", "&")
    return " ".join(text.split())


def strip_rtf_controls(text: str) -> str:
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d*\s?", " ", text)
    text = text.replace("{", " ").replace("}", " ")
    return " ".join(text.split())


def parse_generic_markup_table(path: Path, strip_fn) -> Tuple[List[float], List[str]]:
    """Parse a table from a markup format (.html/.rtf) using a row-preserving strip.

    Row boundaries (</tr>, <br>, \\par) are split on *before* stripping tags/
    control words, since the strip functions collapse everything to one line.
    """
    raw = read_text(path)
    row_chunks = re.split(r"</tr>|<br\s*/?>|\\par\b", raw, flags=re.IGNORECASE)

    values: List[float] = []
    labels: List[str] = []
    for chunk in row_chunks:
        line = strip_fn(chunk)
        if not line:
            continue
        values.extend(extract_numbers(line))
        label_match = re.match(r"^(.*?)(?=[-+]?\d[\d,]*(?:\.\d+)?)", line)
        if label_match:
            candidate = normalize_label(label_match.group(1))
            if candidate and not candidate.startswith("note"):
                labels.append(label_match.group(1).strip(" :|-"))

    deduped_labels = list(dict.fromkeys([l for l in labels if l]))
    return values, deduped_labels


def parse_output_tables(outputs_dir: Path) -> List[OutputTable]:
    files: List[Path] = []
    for pattern in NUMERIC_TABLE_EXTS:
        files.extend(outputs_dir.glob(pattern))

    pdf_data: dict = {}
    for pdf_path in outputs_dir.glob(PDF_GLOB):
        extracted = extract_pdf_table(pdf_path)
        if extracted is not None:
            pdf_data[pdf_path] = extracted
            files.append(pdf_path)

    files = sorted(files, key=natural_sort_key)

    parsed: List[OutputTable] = []
    for idx, path in enumerate(files, start=1):
        suffix = path.suffix.lower()
        if path in pdf_data:
            numeric_values, row_labels = pdf_data[path]
        elif suffix == ".xlsx":
            numeric_values = parse_output_workbook(path)
            row_labels = parse_output_row_labels(path)
        elif suffix == ".tex":
            numeric_values, row_labels = parse_output_tex(path)
        elif suffix in (".html", ".htm"):
            numeric_values, row_labels = parse_generic_markup_table(path, strip_html_tags)
        elif suffix == ".rtf":
            numeric_values, row_labels = parse_generic_markup_table(path, strip_rtf_controls)
        else:
            numeric_values, row_labels = parse_output_text_table(path)

        parsed.append(
            OutputTable(
                index=idx,
                path=path,
                numeric_values=numeric_values,
                row_labels=row_labels,
            )
        )
    return parsed


def find_visual_review_outputs(outputs_dir: Path) -> List[Path]:
    """Enumerate generated outputs that can't be numerically parsed.

    Figures (.png/.jpg/.jpeg), plus any .pdf whose text layer couldn't be
    extracted (see extract_pdf_table - typically a scanned/rasterized PDF),
    need a human (or Claude vision) to compare them against the rendered
    manuscript pages - see "Figure and visual-output comparison" in SKILL.md.
    These are deliberately excluded from parse_output_tables so they can't
    silently shift the label/index-based pairing used for numeric tables.
    """
    files: List[Path] = []
    for pattern in VISUAL_REVIEW_EXTS:
        files.extend(outputs_dir.glob(pattern))
    for pdf_path in outputs_dir.glob(PDF_GLOB):
        if extract_pdf_table(pdf_path) is None:
            files.append(pdf_path)
    return sorted(files, key=natural_sort_key)


def find_unclassified_outputs(outputs_dir: Path) -> List[Path]:
    """Anything in the outputs folder matched by neither known extension list.

    This is the safety net for exhibit formats we didn't anticipate (.docx,
    .json, an extension-less export, a typo'd extension, ...). Without it, an
    unrecognized file just never appears in the report - this makes that
    impossible by construction: every file is either compared, queued for
    visual review, or listed here as needing manual triage. .pdf is always
    "classified" (into one bucket or the other via extract_pdf_table), so it
    never lands here.
    """
    classified = set()
    for pattern in NUMERIC_TABLE_EXTS + VISUAL_REVIEW_EXTS + (PDF_GLOB,):
        classified.update(outputs_dir.glob(pattern))
    files = [p for p in outputs_dir.iterdir() if p.is_file() and p not in classified]
    return sorted(files, key=natural_sort_key)


NOTES_LINE_RE = re.compile(r"^\|?\s*notes?\s*:", re.IGNORECASE)
CAPTION_LINE_RE = re.compile(r"^\|?\s*table\s*:", re.IGNORECASE)
GRID_BOUNDARY_RE = re.compile(r"^\+=+\+")


def parse_output_text_table(path: Path) -> Tuple[List[float], List[str]]:
    """Parse a pandoc/modelsummary-style grid text table.

    These exports always end with a "Table: Table N: <caption>" line outside
    the box, and a footer section (a "Notes:" block, or a column-legend line
    like "(1) Micro (1-9), (2) Small (10-19), ..."). Both are prose, not data,
    but a blind per-line number scan picks up "17 sectors", "1%", "5%", "10%",
    range labels like "1-9"/"10-19", and the table number from the caption
    line as if they were real values - inflating the value count (breaking
    the exact-match check even when every real number matches) or, worse,
    smuggling in numbers that then misalign every subsequent positional
    comparison against the manuscript.

    The footer section is structurally delimited: this grid format always has
    exactly one "+===+" (double-line) separator marking the end of the header
    and, when a footer is present, one right after the last data row and one
    closing the footer - i.e. the *last two* such separators bracket the
    footer. That's a more robust signal than matching specific footer wording.
    """
    content = read_text(path)
    lines = [" ".join(raw.split()) for raw in content.splitlines()]

    boundary_idxs = [i for i, line in enumerate(lines) if GRID_BOUNDARY_RE.match(line)]
    footer_start, footer_end = (None, None)
    if len(boundary_idxs) >= 2:
        footer_start, footer_end = boundary_idxs[-2], boundary_idxs[-1]

    values: List[float] = []
    labels: List[str] = []
    in_notes = False

    for i, line in enumerate(lines):
        if not line:
            continue
        if CAPTION_LINE_RE.match(line):
            continue
        if footer_start is not None and footer_start < i < footer_end:
            continue
        if NOTES_LINE_RE.match(line):
            in_notes = True
            continue
        if in_notes:
            continue

        values.extend(extract_numbers(line))

        # Try to capture a leading text label before the first numeric token.
        label_match = re.match(r"^(.*?)(?=[-+]?\d[\d,]*(?:\.\d+)?)", line)
        if label_match:
            candidate = normalize_label(label_match.group(1))
            if candidate and not candidate.startswith("note"):
                labels.append(label_match.group(1).strip(" :|-"))

    deduped_labels = list(dict.fromkeys([l for l in labels if l]))
    return values, deduped_labels


def find_missing_rows(paper_labels: List[str], output_labels: List[str]) -> List[str]:
    output_norm = {normalize_label(label) for label in output_labels}
    output_canonical = {canonicalize_label(label) for label in output_labels}

    # Some generated regression summaries report Countries/Cyber_Incidents but omit
    # FE declaration lines that are often added in manuscript formatting.
    optional_rows: set[str] = set()
    if "countries" in output_canonical and "cyber incidents" in output_canonical:
        optional_rows.update({"country sector fe", "country year fe"})

    missing: List[str] = []
    for label in paper_labels:
        norm = normalize_label(label)
        if not norm:
            continue
        canon = canonicalize_label(label)
        if canon in optional_rows:
            continue
        if norm in output_norm or canon in output_canonical:
            continue

        # Fallback fuzzy token inclusion for close variants.
        tokens = set(canon.split())
        matched = False
        for out in output_canonical:
            out_tokens = set(out.split())
            if tokens and tokens.issubset(out_tokens):
                matched = True
                break
            if out_tokens and out_tokens.issubset(tokens):
                matched = True
                break

        if not matched:
            missing.append(label)
    return missing


def compute_match_rate(a: List[float], b: List[float], tolerance: float) -> float:
    if not a and not b:
        return 1.0
    if not a or not b:
        return 0.0

    n = min(len(a), len(b))
    same = 0
    for i in range(n):
        if abs(a[i] - b[i]) <= tolerance:
            same += 1

    return same / max(len(a), len(b))


def is_exact_match(a: List[float], b: List[float], tolerance: float) -> bool:
    if not a or not b:
        return False
    if len(a) != len(b):
        return False
    return all(abs(a[i] - b[i]) <= tolerance for i in range(len(a)))


TABLE_LABEL_RE = re.compile(r"table[\s_]*([a-z]?\d+[a-z]?)", re.IGNORECASE)

# Many Stata/R export conventions name appendix table files after the label
# alone (e.g. "A1 attrition_training_followup_2021.tex", "A9 results.tex")
# with no literal "table" word in the filename at all. TABLE_LABEL_RE alone
# would silently fail to tokenize these, leaving every appendix output
# unpaired even though the manuscript side resolves "A1".."A9" fine via
# \label{tab:A1}. Only match at the very start of the filename stem so this
# doesn't misfire on numbers buried elsewhere (e.g. a trailing "_2025").
FILENAME_LEADING_LABEL_RE = re.compile(r"^\s*([a-z]\d+[a-z]?|\d+[a-z]?)(?=[\s_]|$)", re.IGNORECASE)


def extract_table_token(text: str) -> Optional[str]:
    """Pull the table number/letter token (e.g. '1', 'A2') out of a caption or filename."""
    match = TABLE_LABEL_RE.search(text)
    if not match:
        return None
    return match.group(1).upper()


def extract_output_filename_token(stem: str) -> Optional[str]:
    """Pull the table token out of a generated-output filename stem.

    Tries the "table N" style first (table_1_descriptive_table.tex), then
    falls back to a leading bare label (A1 attrition_training_followup_2021.tex).
    """
    token = extract_table_token(stem)
    if token:
        return token
    match = FILENAME_LEADING_LABEL_RE.match(stem)
    if not match:
        return None
    return match.group(1).upper()


def format_table_name(label_token: Optional[str], table_index: int) -> str:
    if label_token:
        return f"Table {label_token}"
    return f"Table {table_index}"


def is_annex_table(table_name: str) -> bool:
    return bool(re.match(r"^table\s+a\d+", table_name, flags=re.IGNORECASE))


def pair_tables(
    paper_tables: List[PaperTable], output_tables: List[OutputTable]
) -> List[Tuple[Optional[PaperTable], Optional[OutputTable]]]:
    """Pair each manuscript table with its generated output.

    Prefer matching by the table label encoded in the output filename
    (table3.txt -> "Table 3", tableA1.txt -> "Table A1") over raw positional
    order. Positional/natural-sort pairing alone breaks whenever a non-table
    numeric file (e.g. a figure's underlying data, Figure1(Data).csv) sits in
    the outputs folder ahead of the real table exports and silently shifts
    every subsequent pairing by one.

    Positional fallback is used *only* for paper tables that have no
    resolvable label at all (a rare parsing edge case). A paper table with a
    known label (the normal case) that finds no matching output is left
    unmatched rather than being handed an arbitrary unlabeled file - e.g. if
    "Table 1"'s real output is table1.pdf (routed to the visual-review
    bucket, not here), it must not get paired with an unrelated leftover like
    a figure's underlying-data CSV just because both lack a numeric partner.
    """
    labeled_outputs: dict = {}
    unlabeled_outputs: List[OutputTable] = []
    for out in output_tables:
        token = extract_output_filename_token(out.path.stem)
        if token and token not in labeled_outputs:
            labeled_outputs[token] = out
        else:
            unlabeled_outputs.append(out)

    unlabeled_iter = iter(unlabeled_outputs)
    pairs: List[Tuple[Optional[PaperTable], Optional[OutputTable]]] = []
    for paper in paper_tables:
        if paper.label_token:
            out = labeled_outputs.pop(paper.label_token, None)
        else:
            out = next(unlabeled_iter, None)
        pairs.append((paper, out))

    # Any output table never claimed by a paper table (an extra exhibit not
    # present in the manuscript, or a leftover unlabeled file) still needs to
    # surface in the report rather than being silently dropped.
    used_ids = {id(out) for _, out in pairs if out is not None}
    for out in output_tables:
        if id(out) not in used_ids:
            pairs.append((None, out))
            used_ids.add(id(out))

    return pairs


def build_rows(
    paper_tables: List[PaperTable],
    output_tables: List[OutputTable],
    tolerance: float,
) -> List[dict]:
    rows: List[dict] = []
    pairs = pair_tables(paper_tables, output_tables)

    for i, (paper, out) in enumerate(pairs):
        p_vals = paper.numeric_values if paper else []
        o_vals = out.numeric_values if out else []
        rate = compute_match_rate(p_vals, o_vals, tolerance)
        has_data = bool(p_vals and o_vals)
        missing_rows = find_missing_rows(
            paper.row_labels if paper else [],
            out.row_labels if out else [],
        )
        if paper:
            table_name = format_table_name(paper.label_token, i + 1)
        else:
            # No paper table claimed this output - it's an extra/orphan file
            # (e.g. Figure1.xlsx, the raw data behind a figure, not a table
            # export). Falling back to "Table {i+1}" here fabricates a table
            # number that doesn't exist in the manuscript (an 8-table paper
            # would get a phantom "Table 9") - name it after the file instead.
            out_token = extract_output_filename_token(out.path.stem) if out else None
            if out_token:
                table_name = f"Table {out_token}"
            elif out:
                table_name = f"Extra output ({out.path.name})"
            else:
                table_name = f"Extra output {i + 1}"

        reason = ""
        if paper is None:
            status = "Not verified"
            reason = "generated output has no corresponding manuscript table"
        elif out is None:
            status = "Not verified"
            reason = "generated output table missing"
        elif not has_data:
            status = "Not verified"
            reason = "insufficient numeric data for comparison"
        elif missing_rows:
            status = "Not verified"
            reason = "missing manuscript rows in generated output"
        elif is_exact_match(p_vals, o_vals, tolerance):
            status = "Verified"
            reason = "exact numeric reproduction"
        else:
            status = "Not verified"
            reason = "numeric differences"

        rows.append(
            {
                "table_index": i + 1,
                "table_name": table_name,
                "paper_caption": paper.caption if paper else "(missing)",
                "output_file": out.path.name if out else "(missing)",
                "paper_value_count": len(p_vals),
                "output_value_count": len(o_vals),
                "match_rate": round(rate, 4),
                "status": status,
                "reason": reason,
                "paper_row_count": len(paper.row_labels) if paper else 0,
                "output_row_count": len(out.row_labels) if out else 0,
                "missing_rows": missing_rows,
                "section": "Annex" if is_annex_table(table_name) else "Main",
            }
        )

    return rows


def build_reviewer_summary(rows: List[dict]) -> List[str]:
    not_verified = [r for r in rows if r["status"] != "Verified"]
    if not not_verified:
        return ["All compared manuscript tables are Verified against Step 2 generated outputs."]

    lines = ["Some compared manuscript tables are not verified against Step 2 generated outputs."]
    for row in not_verified:
        lines.append(f"{row['table_name']}: {row['reason']}")
    return lines


def write_markdown(
    report_path: Path,
    package_id: str,
    rows: List[dict],
    step2_failed: bool,
    step2_reason: str,
    visual_outputs: Optional[List[Path]] = None,
    unclassified_outputs: Optional[List[Path]] = None,
) -> None:
    visual_outputs = visual_outputs or []
    unclassified_outputs = unclassified_outputs or []
    if step2_failed:
        lines = [
            f"# Table Verification Report - {package_id}",
            "",
            "Package execution status: FAILED in Step 2.",
            "Comparison status: Not performed.",
            f"Reason: {step2_reason or 'Step 2 run log indicates execution failure.'}",
            "",
            "Result: Could not verify manuscript findings because the code did not complete successfully.",
        ]
        report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        return

    verified = sum(1 for r in rows if r["status"] == "Verified")
    not_verified = sum(1 for r in rows if r["status"] != "Verified")
    reviewer_summary = build_reviewer_summary(rows)
    main_rows = [r for r in rows if r["section"] == "Main"]
    annex_rows = [r for r in rows if r["section"] == "Annex"]

    lines = [
        f"# Table Verification Report - {package_id}",
        "",
        "Package execution status: SUCCESS in Step 2.",
        "Comparison scope: manuscript tables vs Step 2 generated table outputs.",
        "",
        "## Summary",
        f"- Verified: {verified}",
        f"- Not verified: {not_verified}",
        "",
        "## Reviewer-ready summary",
        *[f"- {line}" for line in reviewer_summary],
    ]

    if main_rows:
        lines.extend(["", "Results in the Main Section of the Paper"])
        for row in main_rows:
            if row["status"] == "Verified":
                lines.append(
                    f"- {row['table_name']} Verified. Verified against {row['output_file']}."
                )
            else:
                lines.append(
                    f"- {row['table_name']} Not verified. Compared against {row['output_file']}. Reason: {row['reason']}."
                )

    if annex_rows:
        lines.extend(["", "Results in the Annex"])
        for row in annex_rows:
            if row["status"] == "Verified":
                lines.append(
                    f"- {row['table_name']} Verified. Verified against {row['output_file']}."
                )
            else:
                lines.append(
                    f"- {row['table_name']} Not verified. Compared against {row['output_file']}. Reason: {row['reason']}."
                )

    if visual_outputs:
        lines.extend([
            "",
            "## Figures and PDF/image outputs - requires manual visual review",
            "These are not numerically parsed. Compare each against the rendered manuscript",
            "pages using Claude vision (see SKILL.md, 'Figure and visual-output comparison'),",
            "then add a Verified / Not verified line for it above.",
        ])
        for path in visual_outputs:
            lines.append(f"- {path.name}")

    if unclassified_outputs:
        lines.extend([
            "",
            "## Unclassified outputs - unrecognized format, needs manual triage",
            "These files were found in the outputs folder but don't match any table",
            "or figure format this script knows how to handle. Confirm what each one is",
            "and either compare it manually or note why it doesn't need an exhibit verdict.",
        ])
        for path in unclassified_outputs:
            lines.append(f"- {path.name}")

    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main(argv: Optional[Iterable[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Compare paper tables to generated outputs")
    parser.add_argument("--package-id", required=True)
    parser.add_argument("--paper-tex", required=True, type=Path)
    parser.add_argument("--outputs-dir", required=True, type=Path)
    parser.add_argument("--report-out", required=True, type=Path)
    parser.add_argument("--json-out", required=True, type=Path)
    parser.add_argument("--step2-log", type=Path)
    parser.add_argument(
        "--step2-status",
        choices=["success", "failed", "unknown"],
        default="unknown",
    )
    parser.add_argument("--tolerance", type=float, default=1e-6)

    args = parser.parse_args(argv)

    step2_failed = args.step2_status == "failed"
    step2_reason = ""
    if args.step2_status == "unknown" and args.step2_log:
        step2_failed, step2_reason = detect_step2_failure(args.step2_log)

    if step2_failed:
        rows: List[dict] = []
        paper_tables = []
        output_tables = []
        visual_outputs: List[Path] = []
        unclassified_outputs: List[Path] = []
    else:
        if not args.paper_tex.exists():
            print(f"ERROR: paper tex not found: {args.paper_tex}", file=sys.stderr)
            return 2

        if not args.outputs_dir.exists():
            print(f"ERROR: outputs dir not found: {args.outputs_dir}", file=sys.stderr)
            return 2

        paper_tables = parse_paper_tables(args.paper_tex)
        output_tables = parse_output_tables(args.outputs_dir)
        rows = build_rows(paper_tables, output_tables, tolerance=args.tolerance)
        visual_outputs = find_visual_review_outputs(args.outputs_dir)
        unclassified_outputs = find_unclassified_outputs(args.outputs_dir)

    args.report_out.parent.mkdir(parents=True, exist_ok=True)
    args.json_out.parent.mkdir(parents=True, exist_ok=True)

    write_markdown(
        args.report_out,
        args.package_id,
        rows,
        step2_failed=step2_failed,
        step2_reason=step2_reason,
        visual_outputs=visual_outputs,
        unclassified_outputs=unclassified_outputs,
    )
    if step2_failed:
        reviewer_summary = [
            "Comparison was not performed because Step 2 code execution failed.",
            step2_reason or "Step 2 run log indicates execution failure.",
        ]
    else:
        reviewer_summary = build_reviewer_summary(rows)

    payload = {
        "package_id": args.package_id,
        "paper_tex": str(args.paper_tex),
        "outputs_dir": str(args.outputs_dir),
        "step2_status": "failed" if step2_failed else "success",
        "step2_reason": step2_reason,
        "table_count_paper": len(paper_tables),
        "table_count_outputs": len(output_tables),
        "reviewer_summary": reviewer_summary,
        "rows": rows,
        "visual_review_outputs": [p.name for p in visual_outputs],
        "unclassified_outputs": [p.name for p in unclassified_outputs],
    }
    args.json_out.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    print(f"Wrote report: {args.report_out}")
    print(f"Wrote json:   {args.json_out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
